"""
tests/test_bug_fixes_verification.py
────────────────────────────────────
Comprehensive automated test suite covering all 16 required verification
scenarios from Section G of the Master Refactoring Specification.

Uses temporary in-memory workbooks and mock-injected timestamps for
deterministic, non-destructive execution.
"""

import time
import pytest
import openpyxl
from unittest.mock import MagicMock

from config import settings
from core.models import (
    LatchState,
    MeasurementRecord,
    SaveStatus,
    StabilityStatus,
    SystemState,
    WriteMode,
)
from core.exceptions import WorkbookLockedError
from infrastructure.excel_repository import ExcelRepository
from services.excel_service import ExcelService
from services.stability_detector import StabilityDetector
from services.auto_save_service import AutoSaveService
from controllers.measurement_controller import MeasurementController


def make_clean_workbook(sheet_name="FLAP_MEASUREMENT"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    # Populate rows 1-10 with header/metadata simulation
    for r in range(1, 11):
        ws[f"A{r}"] = f"Header A{r}"
        ws[f"G{r}"] = f"Meta G{r}" if r < 11 else None
    return wb, ws


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 1: Memilih dan mengedit row pertama, row tengah, dan row terakhir
# ══════════════════════════════════════════════════════════════════════
def test_scenario_01_select_and_edit_first_middle_last_row():
    wb, ws = make_clean_workbook()
    repo = ExcelRepository()
    repo.inject(wb, "test_scen1.xlsx")
    svc = ExcelService(repo)
    svc.activate_sheet("FLAP_MEASUREMENT")

    # Populate rows 11, 12, 13
    svc.append_measurement(100.0)  # row 11 (first)
    svc.append_measurement(200.0)  # row 12 (middle)
    svc.append_measurement(300.0)  # row 13 (last)

    assert svc.next_append_row == 14

    # 1. Select and edit first row (row 11)
    svc.select_row(history_index=0, excel_row=11)
    assert svc.selected_excel_row == 11
    svc.update_record(11, 268.70)
    assert ws["G11"].value == 268.70
    assert ws["G12"].value == 200.0
    assert ws["G13"].value == 300.0
    assert svc.next_append_row == 14  # Unchanged!

    # 2. Select and edit middle row (row 12)
    svc.select_row(history_index=1, excel_row=12)
    assert svc.selected_excel_row == 12
    svc.update_record(12, 268.42)
    assert ws["G11"].value == 268.70
    assert ws["G12"].value == 268.42
    assert ws["G13"].value == 300.0
    assert svc.next_append_row == 14  # Unchanged!

    # 3. Select and edit last row (row 13)
    svc.select_row(history_index=2, excel_row=13)
    assert svc.selected_excel_row == 13
    svc.update_record(13, 271.34)
    assert ws["G11"].value == 268.70
    assert ws["G12"].value == 268.42
    assert ws["G13"].value == 271.34
    assert svc.next_append_row == 14  # Unchanged!


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 2: Berpindah row berulang kali tanpa mengubah target penulisan secara tidak sengaja
# ══════════════════════════════════════════════════════════════════════
def test_scenario_02_switch_rows_repeatedly_without_changing_target_append():
    wb, ws = make_clean_workbook()
    repo = ExcelRepository()
    repo.inject(wb, "test_scen2.xlsx")
    svc = ExcelService(repo)
    svc.activate_sheet("FLAP_MEASUREMENT")

    svc.append_measurement(10.0)  # row 11
    svc.append_measurement(20.0)  # row 12
    svc.append_measurement(30.0)  # row 13
    assert svc.next_append_row == 14

    # Switch selection back and forth
    for _ in range(5):
        svc.select_row(0, 11)
        assert svc.selected_excel_row == 11
        assert svc.next_append_row == 14
        assert svc.write_mode == WriteMode.APPEND

        svc.select_row(2, 13)
        assert svc.selected_excel_row == 13
        assert svc.next_append_row == 14
        assert svc.write_mode == WriteMode.APPEND

        svc.select_row(1, 12)
        assert svc.selected_excel_row == 12
        assert svc.next_append_row == 14
        assert svc.write_mode == WriteMode.APPEND


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 3: Append record baru setelah edit row
# ══════════════════════════════════════════════════════════════════════
def test_scenario_03_append_new_record_after_edit_row():
    wb, ws = make_clean_workbook()
    repo = ExcelRepository()
    repo.inject(wb, "test_scen3.xlsx")
    svc = ExcelService(repo)
    svc.activate_sheet("FLAP_MEASUREMENT")

    svc.append_measurement(10.0)  # row 11
    svc.append_measurement(20.0)  # row 12
    assert svc.next_append_row == 13

    # Edit row 11
    svc.select_row(0, 11)
    svc.update_record(11, 15.0)
    assert ws["G11"].value == 15.0
    assert svc.next_append_row == 13

    # Now append new measurement
    res = svc.append_measurement(30.0)
    assert res.excel_row == 13
    assert ws["G13"].value == 30.0
    assert ws["G11"].value == 15.0  # row 11 was not overwritten
    assert ws["G12"].value == 20.0  # row 12 was not overwritten
    assert svc.next_append_row == 14


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 4: Pause saat sensor bergerak
# ══════════════════════════════════════════════════════════════════════
def test_scenario_04_pause_while_sensor_moving():
    wb, ws = make_clean_workbook()
    repo = ExcelRepository()
    repo.inject(wb, "test_scen4.xlsx")
    svc = ExcelService(repo)
    svc.activate_sheet("FLAP_MEASUREMENT")

    detector = StabilityDetector(movement_threshold_deg=0.08)
    auto_save = AutoSaveService(excel_service=svc, stability_detector=detector)
    auto_save.enable()
    assert auto_save.state == SystemState.RUNNING

    # Sensor in motion
    auto_save.feed(10.0, monotonic_time=1.0)
    auto_save.feed(12.0, monotonic_time=1.3)
    assert detector.status == StabilityStatus.MOVING

    # Operator pauses
    auto_save.pause()
    assert auto_save.state == SystemState.PAUSED

    # Feed while paused
    res = auto_save.feed(14.0, monotonic_time=1.6)
    assert res is None
    assert len(svc.history) == 0


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 5: Pause ketika kandidat stabil sedang terbentuk
# ══════════════════════════════════════════════════════════════════════
def test_scenario_05_pause_while_stable_candidate_forming():
    wb, ws = make_clean_workbook()
    repo = ExcelRepository()
    repo.inject(wb, "test_scen5.xlsx")
    svc = ExcelService(repo)
    svc.activate_sheet("FLAP_MEASUREMENT")

    detector = StabilityDetector(min_samples=8, min_duration_sec=1.0)
    auto_save = AutoSaveService(excel_service=svc, stability_detector=detector)
    auto_save.enable()

    # Feed 6 samples (accumulating in SETTLING)
    t = 0.0
    for _ in range(6):
        auto_save.feed(268.70, monotonic_time=t)
        t += 0.15
    assert detector.status == StabilityStatus.SETTLING

    # Pause pressed before 8 samples / 1.0s reached
    auto_save.pause()
    assert auto_save.state == SystemState.PAUSED
    assert detector.status == StabilityStatus.WAITING_FOR_DATA  # Reset cleanly!

    # Feed remaining samples while paused
    for _ in range(10):
        res = auto_save.feed(268.70, monotonic_time=t)
        assert res is None
        t += 0.15

    assert len(svc.history) == 0  # No auto-save occurred!


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 6: Pause ketika terdapat job Auto-Save di queue / pending
# ══════════════════════════════════════════════════════════════════════
def test_scenario_06_pause_with_job_pending():
    mock_excel = MagicMock(spec=ExcelService)
    detector = StabilityDetector()
    auto_save = AutoSaveService(excel_service=mock_excel, stability_detector=detector)
    auto_save.enable()

    # Pause is called
    auto_save.pause()
    assert auto_save.state == SystemState.PAUSED

    # Direct commit attempt with stale generation
    res = auto_save._attempt_save(angle=268.70, is_simulator=False, generation_id=0)
    assert res is None
    mock_excel.append_measurement.assert_not_called()


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 7: Resume setelah Pause
# ══════════════════════════════════════════════════════════════════════
def test_scenario_07_resume_after_pause():
    wb, ws = make_clean_workbook()
    repo = ExcelRepository()
    repo.inject(wb, "test_scen7.xlsx")
    svc = ExcelService(repo)
    svc.activate_sheet("FLAP_MEASUREMENT")

    detector = StabilityDetector(min_samples=5, min_duration_sec=0.5, tolerance_deg=0.05)
    auto_save = AutoSaveService(excel_service=svc, stability_detector=detector, sync_save=True)
    auto_save.enable()
    auto_save.pause()
    assert auto_save.state == SystemState.PAUSED

    auto_save.resume()
    assert auto_save.state == SystemState.RUNNING

    # After resume, fresh samples require motion cycle before saving
    t = 100.0
    for _ in range(5):
        auto_save.feed(20.0, monotonic_time=t)
        t += 0.15
    for val in [21.0, 22.0, 23.0, 24.0]:
        auto_save.feed(val, monotonic_time=t)
        t += 0.15
    for _ in range(15):
        auto_save.feed(25.0, monotonic_time=t)
        t += 0.15

    assert len(svc.history) == 1
    assert ws["G11"].value == 25.0


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 8: Memastikan event dari sesi sebelum Pause tidak menulis data
# ══════════════════════════════════════════════════════════════════════
def test_scenario_08_stale_pre_pause_events_do_not_write_data():
    wb, ws = make_clean_workbook()
    repo = ExcelRepository()
    repo.inject(wb, "test_scen8.xlsx")
    svc = ExcelService(repo)
    svc.activate_sheet("FLAP_MEASUREMENT")

    detector = StabilityDetector(min_samples=8, min_duration_sec=1.0)
    auto_save = AutoSaveService(excel_service=svc, stability_detector=detector)
    auto_save.enable()

    # Feed 7 samples pre-pause
    t = 0.0
    for _ in range(7):
        auto_save.feed(268.70, monotonic_time=t)
        t += 0.15

    # Operator pauses
    auto_save.pause()

    # Operator resumes
    auto_save.resume()

    # A single sample arrives (which would have been sample #8 pre-pause)
    res = auto_save.feed(268.70, monotonic_time=t)
    # MUST NOT fire save because pre-pause window was discarded!
    assert res is None
    assert len(svc.history) == 0


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 9: Pembacaan 268.70 → 268.42 dan gerakan lanjutan
# ══════════════════════════════════════════════════════════════════════
def test_scenario_09_reading_268_70_to_268_42_and_further():
    wb, ws = make_clean_workbook()
    repo = ExcelRepository()
    repo.inject(wb, "test_scen9.xlsx")
    svc = ExcelService(repo)
    svc.activate_sheet("FLAP_MEASUREMENT")

    detector = StabilityDetector(
        tolerance_deg=0.05,
        min_samples=5,
        min_duration_sec=0.5,
        rearm_threshold_deg=0.18,
        movement_threshold_deg=0.08,
    )
    auto_save = AutoSaveService(excel_service=svc, stability_detector=detector, sync_save=True)
    auto_save.enable()

    t = 0.0
    # Baseline
    for _ in range(5):
        auto_save.feed(268.00, monotonic_time=t)
        t += 0.15

    # 1. First position: 268.70
    for val in [268.20, 268.40, 268.60]:
        auto_save.feed(val, monotonic_time=t)
        t += 0.15
    for _ in range(12):
        auto_save.feed(268.70, monotonic_time=t)
        t += 0.15
    assert len(svc.history) == 1
    assert ws["G11"].value == 268.70

    # 2. Transition: moves to 268.42 (delta = 0.28° > 0.18°)
    for val in [268.55, 268.48]:
        auto_save.feed(val, monotonic_time=t)
        t += 0.15
    for _ in range(12):
        auto_save.feed(268.42, monotonic_time=t)
        t += 0.15
    assert len(svc.history) == 2
    assert ws["G12"].value == 268.42

    # 3. Transition: moves to 271.34
    for val in [269.00, 270.00, 271.00]:
        auto_save.feed(val, monotonic_time=t)
        t += 0.15
    for _ in range(12):
        auto_save.feed(271.34, monotonic_time=t)
        t += 0.15
    assert len(svc.history) == 3
    assert ws["G13"].value == 271.34


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 10: Sensor berhenti pada posisi dengan noise kecil
# ══════════════════════════════════════════════════════════════════════
def test_scenario_10_sensor_stops_with_small_noise():
    wb, ws = make_clean_workbook()
    repo = ExcelRepository()
    repo.inject(wb, "test_scen10.xlsx")
    svc = ExcelService(repo)
    svc.activate_sheet("FLAP_MEASUREMENT")

    detector = StabilityDetector(
        tolerance_deg=0.05,
        min_samples=5,
        min_duration_sec=0.5,
    )
    auto_save = AutoSaveService(excel_service=svc, stability_detector=detector, sync_save=True)
    auto_save.enable()

    t = 0.0
    # Baseline
    for _ in range(5):
        auto_save.feed(265.00, monotonic_time=t)
        t += 0.15

    # Move to 268.70
    for val in [266.00, 267.00, 268.00]:
        auto_save.feed(val, monotonic_time=t)
        t += 0.15

    # Small synchro jitter around 268.70
    jitter = [268.70, 268.71, 268.69, 268.70, 268.72, 268.70] * 2
    for a in jitter:
        auto_save.feed(a, monotonic_time=t)
        t += 0.15

    assert len(svc.history) == 1
    # Median representative value is approximately 268.70
    assert abs(ws["G11"].value - 268.70) <= 0.05


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 11: Sensor terus bergerak sehingga tidak boleh dianggap stabil
# ══════════════════════════════════════════════════════════════════════
def test_scenario_11_sensor_continuously_moving_never_stable():
    detector = StabilityDetector(
        tolerance_deg=0.05,
        min_samples=5,
        min_duration_sec=0.5,
        movement_threshold_deg=0.08,
    )
    t = 0.0
    angle = 10.0
    for _ in range(30):
        res = detector.update(angle, monotonic_time=t)
        assert res.status != StabilityStatus.STABLE
        assert res.just_became_stable is False
        angle += 0.15  # continuously moving
        t += 0.1


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 12: Sensor stabil lalu berpindah ke posisi baru
# ══════════════════════════════════════════════════════════════════════
def test_scenario_12_sensor_stable_then_moves_to_new_position():
    wb, ws = make_clean_workbook()
    repo = ExcelRepository()
    repo.inject(wb, "test_scen12.xlsx")
    svc = ExcelService(repo)
    svc.activate_sheet("FLAP_MEASUREMENT")

    detector = StabilityDetector(min_samples=5, min_duration_sec=0.5, rearm_threshold_deg=0.20)
    auto_save = AutoSaveService(excel_service=svc, stability_detector=detector, sync_save=True)
    auto_save.enable()

    t = 0.0
    # Baseline
    for _ in range(5):
        auto_save.feed(5.0, monotonic_time=t)
        t += 0.15

    # Move to 10.0
    for val in [6.0, 7.5, 9.0]:
        auto_save.feed(val, monotonic_time=t)
        t += 0.15

    # First position (10.0)
    for _ in range(12):
        auto_save.feed(10.0, monotonic_time=t)
        t += 0.15
    assert len(svc.history) == 1
    assert auto_save.latch_state == LatchState.LATCHED

    # Stationary holding — no duplicate saves
    for _ in range(10):
        auto_save.feed(10.0, monotonic_time=t)
        t += 0.15
    assert len(svc.history) == 1

    # Move to new position (15.0)
    for val in [11.5, 13.0, 14.2]:
        auto_save.feed(val, monotonic_time=t)
        t += 0.15
    for _ in range(12):
        auto_save.feed(15.0, monotonic_time=t)
        t += 0.15
    assert len(svc.history) == 2
    assert ws["G12"].value == 15.0


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 13: Excel gagal ditulis lalu retry
# ══════════════════════════════════════════════════════════════════════
def test_scenario_13_excel_write_failure_and_retry():
    wb, ws = make_clean_workbook()
    repo = ExcelRepository()
    repo.inject(wb, "test_scen13.xlsx")

    # Simulate 2 lock errors then success
    call_counts = {"save": 0}
    orig_save = repo.save

    def flakey_save():
        call_counts["save"] += 1
        if call_counts["save"] < 3:
            raise WorkbookLockedError("File temporarily locked")
        orig_save()

    repo.save = flakey_save
    repo.write_cell("FLAP_MEASUREMENT", "G", 11, 25.0)
    repo.save_with_retry(max_retries=3, delay_sec=0.01)

    assert call_counts["save"] == 3
    assert ws["G11"].value == 25.0


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 14: Retry tidak menggandakan record
# ══════════════════════════════════════════════════════════════════════
def test_scenario_14_retry_does_not_duplicate_records():
    wb, ws = make_clean_workbook()
    repo = ExcelRepository()
    repo.inject(wb, "test_scen14.xlsx")
    svc = ExcelService(repo)
    svc.activate_sheet("FLAP_MEASUREMENT")

    call_counts = {"count": 0}
    orig_save = repo.save

    def flakey_save():
        call_counts["count"] += 1
        if call_counts["count"] < 2:
            raise WorkbookLockedError("Lock")
        orig_save()

    repo.save = flakey_save

    res = svc.append_measurement(25.0)
    assert res.status == SaveStatus.SAVED
    assert len(svc.history) == 1
    assert len(svc.session_stack) == 1
    assert svc.next_append_row == 12


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 15: Format numeric Excel tetap sesuai dengan kontrak workbook
# ══════════════════════════════════════════════════════════════════════
def test_scenario_15_numeric_format_preserves_contract():
    wb, ws = make_clean_workbook()
    repo = ExcelRepository()
    repo.inject(wb, "test_scen15.xlsx")
    svc = ExcelService(repo)
    svc.activate_sheet("FLAP_MEASUREMENT")

    test_angles = [268.70, 268.42, 271.34, -12.35]
    for a in test_angles:
        svc.append_measurement(a)

    for i, expected in enumerate(test_angles):
        row = 11 + i
        cell = ws[f"G{row}"]
        # Contract 1: Value is Python float
        assert isinstance(cell.value, float)
        assert cell.value == pytest.approx(expected, abs=0.001)
        # Contract 2: Cell type is numeric
        assert cell.data_type == "n"
        # Contract 3: Number format is 0.00
        assert cell.number_format == "0.00"
        # Contract 4: Read back via repository matches float
        assert repo.read_cell("FLAP_MEASUREMENT", "G", row) == cell.value


# ══════════════════════════════════════════════════════════════════════
#  SCENARIO 16: Worksheet dan struktur baris awal tetap benar (Row 11)
# ══════════════════════════════════════════════════════════════════════
def test_scenario_16_worksheet_and_initial_row_structure():
    wb, ws = make_clean_workbook()
    repo = ExcelRepository()
    repo.inject(wb, "test_scen16.xlsx")
    svc = ExcelService(repo)
    svc.activate_sheet("FLAP_MEASUREMENT")

    # Initial append must start at Row 11
    assert svc.next_append_row == 11

    # Rows 1-10 headers must remain completely intact
    for r in range(1, 11):
        assert ws[f"A{r}"].value == f"Header A{r}"
        if r < 11:
            assert ws[f"G{r}"].value == f"Meta G{r}"

    # First write lands on row 11
    res = svc.append_measurement(18.5)
    assert res.excel_row == 11
    assert ws["G11"].value == 18.5
    assert svc.next_append_row == 12

    # Headers 1-10 still untouched
    for r in range(1, 11):
        assert ws[f"A{r}"].value == f"Header A{r}"
