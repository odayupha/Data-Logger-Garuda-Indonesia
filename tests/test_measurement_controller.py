"""
tests/test_measurement_controller.py
─────────────────────────────────────
Unit tests for MeasurementController facade.
"""

import pytest
import openpyxl
from controllers.measurement_controller import MeasurementController
from core.models import SaveStatus, StabilityStatus, LatchState


def create_in_memory_controller(sheet_name="FLAP_MEASUREMENT"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    auto_saves = []
    stabilities = []
    errors = []

    ctrl = MeasurementController(
        on_auto_save=auto_saves.append,
        on_stability=stabilities.append,
        on_error=errors.append,
    )
    ctrl.inject_workbook(wb, "test_ctrl_unit.xlsx")
    return ctrl, wb, ws, auto_saves, stabilities, errors


class TestMeasurementControllerWorkbook:
    def test_inject_workbook_and_sheet_selection(self):
        ctrl, wb, ws, _, _, _ = create_in_memory_controller()
        assert ctrl.is_workbook_loaded is True
        assert ctrl.active_sheet == "FLAP_MEASUREMENT"
        assert ctrl.current_row == 11

    def test_get_cell_value(self):
        ctrl, wb, ws, _, _, _ = create_in_memory_controller()
        ws["G11"] = 42.5
        val = ctrl.get_cell_value("FLAP_MEASUREMENT", "G", 11)
        assert val == 42.5


class TestMeasurementControllerManualSave:
    def test_manual_save(self):
        ctrl, wb, ws, _, _, _ = create_in_memory_controller()
        res = ctrl.save_measurement(25.0)
        assert res.status == SaveStatus.SAVED
        assert res.excel_row == 11
        assert ws["G11"].value == 25.0
        assert ctrl.current_row == 12
        assert ctrl.save_count == 1

    def test_undo_and_edit(self):
        ctrl, wb, ws, _, _, _ = create_in_memory_controller()
        ctrl.save_measurement(25.0)
        ctrl.save_measurement(30.0)
        assert ctrl.save_count == 2

        # Edit record
        edited = ctrl.edit_record(12, 31.5)
        assert edited is not None
        assert edited.angle == 31.5
        assert ws["G12"].value == 31.5

        # Undo last
        undone = ctrl.undo_last()
        assert undone is not None
        assert undone.angle == 31.5
        assert ws["G12"].value is None
        assert ctrl.save_count == 1


class TestMeasurementControllerAutoSave:
    def test_auto_save_enable_disable(self):
        ctrl, _, _, _, _, _ = create_in_memory_controller()
        assert ctrl.auto_save_enabled is False

        ctrl.enable_auto_save()
        assert ctrl.auto_save_enabled is True

        ctrl.disable_auto_save()
        assert ctrl.auto_save_enabled is False

    def test_enable_auto_save_fails_if_workbook_not_loaded(self):
        errors = []
        ctrl = MeasurementController(on_error=errors.append)
        ctrl.enable_auto_save()
        assert ctrl.auto_save_enabled is False
        assert any("no workbook loaded" in e.lower() for e in errors)

    def test_feed_angle_triggers_stability_and_auto_save(self):
        ctrl, wb, ws, auto_saves, stabilities, errors = create_in_memory_controller()
        ctrl.enable_auto_save()

        # Feed stable angles
        for i in range(25):
            ctrl._auto_save.feed(18.5, monotonic_time=i * 0.3)

        assert len(stabilities) > 0
        assert len(auto_saves) == 1
        assert auto_saves[0].status == SaveStatus.SAVED
        assert auto_saves[0].excel_row == 11
        assert ws["G11"].value == 18.5
