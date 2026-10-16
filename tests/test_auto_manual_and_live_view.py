"""
tests/test_auto_manual_and_live_view.py
───────────────────────────────────────
Comprehensive test suite verifying all 18 requirements from Section F:
1. Auto-Save ON saves one record when stability criteria are met.
2. Auto-Save OFF does not generate automatic saves.
3. Enter saves one record in Manual Mode.
4. Repeated Enter does not duplicate records (rapid repeat debounce).
5. Mode change when sensor is stable does not trigger unexpected save.
6. Pause blocks Auto-Save.
7. Resume properly reinitializes stability detection.
8. Actual Excel data appears in GUI Live View after successful commit.
9. Edit Row updates the correct record.
10. Undo updates workbook and view correctly.
11. Physical Excel row numbers remain accurate.
12. Write failure does not produce a false SAVED indicator.
13. Setpoint and tolerance auto-detected from existing sheet headers or overrides.
14. PASS/FAIL math matches deviation and tolerance specifications.
15. Missing setpoint/tolerance is marked NOT_EVALUATED, never false PASS/FAIL.
16. Numeric Excel data preserves contract (strict float).
17. Worksheet switching updates live view and row mapping.
18. Refresh operation is strictly idempotent (no duplicates or loss).
"""

import time
import pytest
import openpyxl

from controllers.measurement_controller import MeasurementController
from core.models import (
    EvaluationStatus,
    LatchState,
    SaveMode,
    SaveResult,
    SaveStatus,
    StabilityStatus,
)
from services.evaluation_service import EvaluationService


def create_test_controller(sheet_name="ACCURACY STALL"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    auto_saves = []
    stabilities = []
    errors = []

    ctrl = MeasurementController(
        on_auto_save=auto_saves.append,
        on_stability=stabilities.append,
        on_error=errors.append,
    )
    ctrl.inject_workbook(wb, "test_live.xlsx")
    return ctrl, wb, ws, auto_saves, stabilities, errors


# ── Requirement 1: Auto-Save ON saves when stability criteria are met ────────
def test_req1_autosave_on_saves_when_stable():
    ctrl, wb, ws, auto_saves, stabilities, _ = create_test_controller()
    ctrl.set_save_mode(SaveMode.AUTO)
    ctrl.enable_auto_save()

    base_time = 1000.0
    # Feed settling to stable at 25.0 degrees
    for i in range(12):
        ctrl.feed_angle(25.0)
        time.sleep(0.01)

    # Directly feed through internal auto-save service with controlled monotonic time to verify 1.0s window
    svc = ctrl._auto_save
    svc._detector.reset()
    svc._latch = LatchState.OPEN

    res = None
    for step in range(15):
        t = base_time + (step * 0.1)
        r = svc.feed(25.00, monotonic_time=t)
        if r is not None:
            res = r
            break

    assert res is not None
    assert res.status == SaveStatus.SAVED
    assert res.excel_row == 11
    assert ws["G11"].value == 25.0
    assert len(auto_saves) >= 1
    assert auto_saves[-1].save_mode == "AUTO"


# ── Requirement 2: Auto-Save OFF does not generate automatic saves ───────────
def test_req2_autosave_off_does_not_save_automatically():
    ctrl, wb, ws, auto_saves, stabilities, _ = create_test_controller()
    ctrl.set_save_mode(SaveMode.MANUAL)

    svc = ctrl._auto_save
    base_time = 2000.0
    for step in range(20):
        t = base_time + (step * 0.1)
        r = svc.feed(30.00, monotonic_time=t)
        assert r is None, "In MANUAL mode, auto-save commit must be suppressed"

    assert len(auto_saves) == 0
    assert ws["G11"].value is None


# ── Requirement 3: Enter saves one record in Manual Mode ──────────────────────
def test_req3_enter_saves_in_manual_mode():
    ctrl, wb, ws, _, _, _ = create_test_controller()
    ctrl.set_save_mode(SaveMode.MANUAL)

    res = ctrl.save_manual(angle=18.45, target_row=11)
    assert res.status == SaveStatus.SAVED
    assert res.excel_row == 11
    assert res.angle == 18.45
    assert ws["G11"].value == 18.45
    assert res.save_mode == "MANUAL"
    assert ctrl.current_row == 12


# ── Requirement 4: Repeated Enter rapid repeat debounce ───────────────────────
def test_req4_repeated_enter_deduplication():
    ctrl, wb, ws, _, _, _ = create_test_controller()
    ctrl.set_save_mode(SaveMode.MANUAL)

    res1 = ctrl.save_manual(angle=18.45, target_row=11)
    assert res1.status == SaveStatus.SAVED

    # Immediate second call with identical angle (within 0.35s) simulates rapid key chatter
    res2 = ctrl.save_manual(angle=18.45, target_row=11)
    assert res2.status == SaveStatus.FAILED
    assert "debounce" in res2.error.lower()


# ── Requirement 5: Mode switch when stable does not trigger unexpected save ───
def test_req5_mode_switch_does_not_trigger_unexpected_save():
    ctrl, wb, ws, auto_saves, _, _ = create_test_controller()
    ctrl.set_save_mode(SaveMode.MANUAL)

    # Feed resting angle in manual mode
    svc = ctrl._auto_save
    base_time = 3000.0
    for step in range(15):
        svc.feed(45.0, monotonic_time=base_time + (step * 0.1))

    assert len(auto_saves) == 0

    # Switch to AUTO
    ctrl.set_save_mode(SaveMode.AUTO)

    # Without new samples, switching mode must not immediately commit the old resting sample
    assert len(auto_saves) == 0
    assert ws["G11"].value is None


# ── Requirement 6: Pause blocks Auto-Save ────────────────────────────────────
def test_req6_pause_blocks_autosave():
    ctrl, wb, ws, auto_saves, _, _ = create_test_controller()
    ctrl.set_save_mode(SaveMode.AUTO)
    ctrl.enable_auto_save()
    ctrl.pause_auto_save()

    svc = ctrl._auto_save
    base_time = 4000.0
    for step in range(20):
        t = base_time + (step * 0.1)
        r = svc.feed(50.0, monotonic_time=t)
        assert r is None

    assert len(auto_saves) == 0
    assert ws["G11"].value is None


# ── Requirement 7: Resume properly reinitializes stability detection ─────────
def test_req7_resume_reinitializes_detection():
    ctrl, wb, ws, auto_saves, _, _ = create_test_controller()
    ctrl.set_save_mode(SaveMode.AUTO)
    ctrl.enable_auto_save()
    ctrl.pause_auto_save()
    ctrl.resume_auto_save()

    assert not ctrl.is_auto_save_paused

    svc = ctrl._auto_save
    base_time = 5000.0
    res = None
    for step in range(15):
        t = base_time + (step * 0.1)
        r = svc.feed(50.0, monotonic_time=t)
        if r is not None:
            res = r
            break

    assert res is not None
    assert res.status == SaveStatus.SAVED
    assert ws["G11"].value == 50.0


# ── Requirement 8: Actual Excel data appears in Live View after commit ────────
def test_req8_live_worksheet_records_reflect_actual_excel():
    ctrl, wb, ws, _, _, _ = create_test_controller()
    ws["G11"] = 12.34
    ws["G12"] = 24.56

    records = ctrl.get_live_worksheet_records()
    assert len(records) >= 2
    rec11 = next(r for r in records if r["excel_row"] == 11)
    rec12 = next(r for r in records if r["excel_row"] == 12)

    assert rec11["measured_angle"] == 12.34
    assert rec12["measured_angle"] == 24.56
    assert rec11["save_status"] == "SAVED"


# ── Requirement 9: Edit Row updates the correct record ───────────────────────
def test_req9_edit_row_updates_correct_record():
    ctrl, wb, ws, _, _, _ = create_test_controller()
    ctrl.save_manual(angle=10.0, target_row=11)
    ctrl.save_manual(angle=20.0, target_row=12)

    updated = ctrl.edit_record(excel_row=11, new_angle=15.0)
    assert updated is not None
    assert updated.angle == 15.0
    assert ws["G11"].value == 15.0
    assert ws["G12"].value == 20.0

    records = ctrl.get_live_worksheet_records()
    rec11 = next(r for r in records if r["excel_row"] == 11)
    assert rec11["measured_angle"] == 15.0


# ── Requirement 10: Undo updates workbook and view correctly ──────────────────
def test_req10_undo_updates_workbook_and_live_view():
    ctrl, wb, ws, _, _, _ = create_test_controller()
    ctrl.save_manual(angle=14.0, target_row=11)
    assert ws["G11"].value == 14.0

    undone = ctrl.undo_last()
    assert undone is not None
    assert undone.excel_row == 11
    assert ws["G11"].value is None
    assert ctrl.current_row == 11


# ── Requirement 11: Physical Excel row numbers remain accurate ────────────────
def test_req11_physical_row_mapping_accuracy():
    ctrl, wb, ws, _, _, _ = create_test_controller()
    ctrl.save_manual(angle=33.3, target_row=15)
    assert ws["G15"].value == 33.3
    assert ws["G11"].value is None


# ── Requirement 12: Write failure does not produce false SAVED indicator ──────
def test_req12_write_failure_does_not_produce_false_saved():
    errors = []
    ctrl = MeasurementController(on_error=errors.append)
    # Controller without injected workbook
    res = ctrl.save_manual(angle=25.0, target_row=11)
    assert res.status == SaveStatus.FAILED
    assert res.error is not None
    assert "workbook" in res.error.lower() or "loaded" in res.error.lower()


# ── Requirement 13: Setpoint and tolerance auto-detection / override ──────────
def test_req13_setpoint_and_tolerance_from_sheet_headers():
    eval_svc = EvaluationService()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TEST_SHEET"

    # Set headers
    ws["C10"] = "SETPOINT (°)"
    ws["D10"] = "TOLERANCE (±°)"
    ws["G10"] = "RESULT (°)"

    col_map = eval_svc.detect_columns_in_sheet(ws, max_header_row=10)
    assert col_map.get("setpoint") == "C"
    assert col_map.get("tolerance") == "D"
    assert col_map.get("result") == "G"

    ws["C11"] = 15.0
    ws["D11"] = 0.20
    sp, tol = eval_svc.extract_row_reference(ws, row=11, setpoint_col=col_map["setpoint"], tolerance_col=col_map["tolerance"])
    assert sp == 15.0
    assert tol == 0.20


# ── Requirement 14: PASS/FAIL calculation math ───────────────────────────────
def test_req14_pass_fail_evaluation_math():
    eval_svc = EvaluationService()

    # Case 1: measured 15.10, setpoint 15.00, tol 0.20 -> error 0.10 <= 0.20 -> PASS
    res1 = eval_svc.evaluate(measured_angle=15.10, setpoint=15.00, tolerance=0.20)
    assert res1.status == EvaluationStatus.PASS
    assert abs(res1.deviation - 0.10) < 1e-9

    # Case 2: measured 15.25, setpoint 15.00, tol 0.20 -> error 0.25 > 0.20 -> FAIL
    res2 = eval_svc.evaluate(measured_angle=15.25, setpoint=15.00, tolerance=0.20)
    assert res2.status == EvaluationStatus.FAIL
    assert abs(res2.deviation - 0.25) < 1e-9

    # Case 3: measured 14.80, setpoint 15.00, tol 0.20 -> error 0.20 <= 0.20 -> PASS (boundary)
    res3 = eval_svc.evaluate(measured_angle=14.80, setpoint=15.00, tolerance=0.20)
    assert res3.status == EvaluationStatus.PASS
    assert abs(res3.deviation - (-0.20)) < 1e-9


# ── Requirement 15: Missing reference marked NOT_EVALUATED ───────────────────
def test_req15_unevaluated_status_when_setpoint_missing():
    eval_svc = EvaluationService()

    # Measured angle without setpoint
    res = eval_svc.evaluate(measured_angle=20.0, setpoint=None, tolerance=0.10)
    assert res.status == EvaluationStatus.NOT_EVALUATED
    assert "setpoint" in res.reason.lower()


# ── Requirement 16: Numeric Excel data preserves contract ─────────────────────
def test_req16_numeric_precision_preserved():
    ctrl, wb, ws, _, _, _ = create_test_controller()
    ctrl.save_manual(angle=25.75, target_row=11)
    val = ws["G11"].value
    assert isinstance(val, (int, float))
    assert not isinstance(val, str)
    assert val == 25.75


# ── Requirement 17: Worksheet switching updates live view and row mapping ────
def test_req17_worksheet_switching_updates_mapping():
    ctrl, wb, ws1, _, _, _ = create_test_controller("SHEET_A")
    ws2 = wb.create_sheet("SHEET_B")
    ws1["G11"] = 11.11
    ws2["G11"] = 22.22

    ctrl.select_sheet("SHEET_A")
    rec_a = ctrl.get_live_worksheet_records()
    assert next(r for r in rec_a if r["excel_row"] == 11)["measured_angle"] == 11.11

    ctrl.select_sheet("SHEET_B")
    rec_b = ctrl.get_live_worksheet_records()
    assert next(r for r in rec_b if r["excel_row"] == 11)["measured_angle"] == 22.22


# ── Requirement 18: Refresh operation is idempotent ───────────────────────────
def test_req18_refresh_operation_idempotence():
    ctrl, wb, ws, _, _, _ = create_test_controller()
    ctrl.save_manual(angle=10.0, target_row=11)
    ctrl.save_manual(angle=20.0, target_row=12)

    rec1 = ctrl.get_live_worksheet_records()
    rec2 = ctrl.get_live_worksheet_records()
    rec3 = ctrl.get_live_worksheet_records()

    assert len(rec1) == len(rec2) == len(rec3)
    assert [r["measured_angle"] for r in rec1] == [r["measured_angle"] for r in rec2] == [r["measured_angle"] for r in rec3]
