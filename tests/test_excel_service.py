"""
tests/test_excel_service.py
───────────────────────────
Unit tests for ExcelService using in-memory ExcelRepository.
"""

import pytest
import openpyxl
from config import settings
from core.exceptions import WorkbookNotLoadedError, WorksheetNotFoundError
from core.models import SaveStatus
from infrastructure.excel_repository import ExcelRepository
from services.excel_service import ExcelService


def create_in_memory_service(sheet_name="FLAP_MEASUREMENT"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    repo = ExcelRepository()
    repo.inject(wb, "memory.xlsx")
    svc = ExcelService(repo)
    svc.activate_sheet(sheet_name)
    return svc, repo, ws


class TestExcelServiceInit:
    def test_initial_state(self):
        repo = ExcelRepository()
        svc = ExcelService(repo)
        assert svc.current_row == settings.EXCEL_START_ROW
        assert svc.active_sheet is None
        assert svc.history == []
        assert svc.session_stack == []
        assert svc.save_count == 0

    def test_current_row_setter_validates(self):
        repo = ExcelRepository()
        svc = ExcelService(repo)
        svc.current_row = 15
        assert svc.current_row == 15
        with pytest.raises(ValueError):
            svc.current_row = 0


class TestSheetActivation:
    def test_activate_sheet_loads_existing_rows(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["G11"] = 12.34
        ws["G12"] = 25.50
        repo = ExcelRepository()
        repo.inject(wb, "test_unit_service.xlsx")

        svc = ExcelService(repo)
        svc.activate_sheet("Sheet1")

        assert svc.active_sheet == "Sheet1"
        assert len(svc.history) == 2
        assert svc.history[0].angle == 12.34
        assert svc.history[1].angle == 25.50
        assert svc.current_row == 13

    def test_activate_unknown_sheet_does_not_crash(self):
        svc, _, _ = create_in_memory_service()
        svc.activate_sheet("UnknownSheet")
        assert svc.active_sheet == "FLAP_MEASUREMENT"


class TestAppendMeasurement:
    def test_append_measurement_success(self):
        callbacks = []
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TestSheet"
        repo = ExcelRepository()
        repo.inject(wb, "test_unit_service.xlsx")

        svc = ExcelService(repo, on_save_complete=callbacks.append)
        svc.activate_sheet("TestSheet")

        result = svc.append_measurement(25.4)

        assert result.status == SaveStatus.SAVED
        assert result.excel_row == 11
        assert ws["G11"].value == 25.4
        assert svc.current_row == 12
        assert len(svc.history) == 1
        assert len(svc.session_stack) == 1
        assert len(callbacks) == 1

    def test_append_measurement_with_target_row(self):
        svc, repo, ws = create_in_memory_service()
        result = svc.append_measurement(30.0, target_row=15)
        assert result.excel_row == 15
        assert ws["G15"].value == 30.0
        assert svc.current_row == 16

    def test_append_measurement_unloaded_raises(self):
        repo = ExcelRepository()
        svc = ExcelService(repo)
        with pytest.raises(WorkbookNotLoadedError):
            svc.append_measurement(20.0)


class TestUndoAndEdit:
    def test_undo_last_record(self):
        svc, repo, ws = create_in_memory_service()
        svc.append_measurement(10.0)
        svc.append_measurement(20.0)
        assert svc.current_row == 13

        removed = svc.undo_last()
        assert removed is not None
        assert removed.angle == 20.0
        assert ws["G12"].value is None
        assert svc.current_row == 12
        assert len(svc.session_stack) == 1
        assert len(svc.history) == 1

    def test_undo_empty_stack_returns_none(self):
        svc, _, _ = create_in_memory_service()
        assert svc.undo_last() is None

    def test_edit_record(self):
        svc, repo, ws = create_in_memory_service()
        svc.append_measurement(10.0)  # G11
        updated = svc.update_record(11, 15.5)
        assert updated is not None
        assert updated.angle == 15.5
        assert ws["G11"].value == 15.5
        assert svc.history[0].angle == 15.5


class TestClearAllData:
    def test_clear_all_data(self):
        svc, repo, ws = create_in_memory_service()
        svc.append_measurement(10.0)
        svc.append_measurement(20.0)
        assert len(svc.history) == 2

        cleared_count = svc.clear_all_data()
        assert cleared_count >= 2
        assert ws["G11"].value is None
        assert ws["G12"].value is None
        assert svc.history == []
        assert svc.session_stack == []
        assert svc.current_row == settings.EXCEL_START_ROW
