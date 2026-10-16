"""
tests/test_excel_repository.py
───────────────────────────────
Unit tests for ExcelRepository using in-memory openpyxl workbooks
(no disk I/O, no temp files required).
"""

import pytest
import openpyxl
from unittest.mock import patch, MagicMock
from infrastructure.excel_repository import ExcelRepository
from core.exceptions import (
    WorkbookNotLoadedError,
    WorksheetNotFoundError,
    WorkbookLockedError,
    WorkbookSaveError,
)


# ── Fixture helpers ────────────────────────────────────────────────────

def make_repo_with_wb(data: dict = None) -> tuple[ExcelRepository, openpyxl.Workbook]:
    """
    Create a repo pre-loaded with an in-memory workbook.

    *data* maps (sheet_name, column, row) → value.
    Default sheet is 'Sheet1'.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    if data:
        for (sheet, col, row), val in data.items():
            if sheet not in wb.sheetnames:
                wb.create_sheet(sheet)
            wb[sheet][f"{col}{row}"] = val

    repo = ExcelRepository()
    repo.inject(wb, "fake_path/test.xlsx")
    return repo, wb


# ══════════════════════════════════════════════════════════════════════
#  Initial state
# ══════════════════════════════════════════════════════════════════════

class TestInitialState:
    def test_not_loaded_initially(self):
        repo = ExcelRepository()
        assert repo.is_loaded is False

    def test_sheet_names_empty_when_not_loaded(self):
        repo = ExcelRepository()
        assert repo.sheet_names == []

    def test_path_is_none_initially(self):
        repo = ExcelRepository()
        assert repo.path is None


# ══════════════════════════════════════════════════════════════════════
#  inject()
# ══════════════════════════════════════════════════════════════════════

class TestInject:
    def test_inject_sets_loaded(self):
        repo, wb = make_repo_with_wb()
        assert repo.is_loaded is True

    def test_inject_exposes_sheet_names(self):
        repo, wb = make_repo_with_wb()
        assert "Sheet1" in repo.sheet_names

    def test_inject_sets_path(self):
        repo, wb = make_repo_with_wb()
        assert repo.path is not None
        assert repo.path.name == "test.xlsx"


# ══════════════════════════════════════════════════════════════════════
#  write_cell + read_cell
# ══════════════════════════════════════════════════════════════════════

class TestWriteReadCell:
    def test_write_and_read_float(self):
        repo, _ = make_repo_with_wb()
        repo.write_cell("Sheet1", "G", 11, 12.34)
        assert repo.read_cell("Sheet1", "G", 11) == 12.34

    def test_write_none_clears_cell(self):
        repo, _ = make_repo_with_wb({("Sheet1", "G", 11): 99.9})
        repo.write_cell("Sheet1", "G", 11, None)
        assert repo.read_cell("Sheet1", "G", 11) is None

    def test_write_raises_if_not_loaded(self):
        repo = ExcelRepository()
        with pytest.raises(WorkbookNotLoadedError):
            repo.write_cell("Sheet1", "G", 11, 1.0)

    def test_write_raises_for_missing_sheet(self):
        repo, _ = make_repo_with_wb()
        with pytest.raises(WorksheetNotFoundError):
            repo.write_cell("NoSuchSheet", "G", 11, 1.0)

    def test_read_raises_if_not_loaded(self):
        repo = ExcelRepository()
        with pytest.raises(WorkbookNotLoadedError):
            repo.read_cell("Sheet1", "G", 11)


# ══════════════════════════════════════════════════════════════════════
#  read_column_range
# ══════════════════════════════════════════════════════════════════════

class TestReadColumnRange:
    def test_reads_contiguous_data(self):
        repo, _ = make_repo_with_wb({
            ("Sheet1", "G", 11): 10.0,
            ("Sheet1", "G", 12): 20.0,
            ("Sheet1", "G", 13): 30.0,
        })
        data = repo.read_column_range("Sheet1", "G", 11)
        assert data == {11: 10.0, 12: 20.0, 13: 30.0}

    def test_stops_at_first_empty(self):
        repo, _ = make_repo_with_wb({
            ("Sheet1", "G", 11): 10.0,
            ("Sheet1", "G", 13): 30.0,   # Gap at row 12
        })
        data = repo.read_column_range("Sheet1", "G", 11, max_consecutive_empty=1)
        assert 11 in data
        assert 13 not in data   # Should stop at gap

    def test_empty_column_returns_empty_dict(self):
        repo, _ = make_repo_with_wb()
        data = repo.read_column_range("Sheet1", "G", 11)
        assert data == {}


# ══════════════════════════════════════════════════════════════════════
#  find_next_empty_row
# ══════════════════════════════════════════════════════════════════════

class TestFindNextEmptyRow:
    def test_returns_start_row_when_column_empty(self):
        repo, _ = make_repo_with_wb()
        row = repo.find_next_empty_row("Sheet1", "G", 11)
        assert row == 11

    def test_returns_row_after_last_data(self):
        repo, _ = make_repo_with_wb({
            ("Sheet1", "G", 11): 1.0,
            ("Sheet1", "G", 12): 2.0,
            ("Sheet1", "G", 13): 3.0,
        })
        row = repo.find_next_empty_row("Sheet1", "G", 11)
        assert row == 14


# ══════════════════════════════════════════════════════════════════════
#  save / save_with_retry
# ══════════════════════════════════════════════════════════════════════

class TestSave:
    def test_save_raises_if_not_loaded(self):
        repo = ExcelRepository()
        with pytest.raises(WorkbookNotLoadedError):
            repo.save()

    def test_save_permission_error_raises_workbook_locked(self):
        repo, wb = make_repo_with_wb()
        with patch.object(wb, "save", side_effect=PermissionError("locked")):
            with pytest.raises(WorkbookLockedError):
                repo.save()

    def test_save_generic_error_raises_workbook_save_error(self):
        repo, wb = make_repo_with_wb()
        with patch.object(wb, "save", side_effect=OSError("disk full")):
            with pytest.raises(WorkbookSaveError):
                repo.save()

    def test_save_with_retry_succeeds_on_second_attempt(self):
        repo, wb = make_repo_with_wb()
        call_count = [0]

        def flaky_save(path):
            call_count[0] += 1
            if call_count[0] == 1:
                raise PermissionError("locked")

        with patch.object(wb, "save", side_effect=flaky_save):
            repo.save_with_retry(max_retries=3, delay_sec=0.0)

        assert call_count[0] == 2

    def test_save_with_retry_raises_after_max_retries(self):
        repo, wb = make_repo_with_wb()
        with patch.object(wb, "save", side_effect=PermissionError("always locked")):
            with pytest.raises(WorkbookLockedError):
                repo.save_with_retry(max_retries=3, delay_sec=0.0)


# ══════════════════════════════════════════════════════════════════════
#  close
# ══════════════════════════════════════════════════════════════════════

class TestClose:
    def test_close_marks_not_loaded(self):
        repo, _ = make_repo_with_wb()
        repo.close()
        assert repo.is_loaded is False
        assert repo.path is None
