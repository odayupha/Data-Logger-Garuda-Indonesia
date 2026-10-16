"""
tests/test_master_ui_excel_integration.py
──────────────────────────────────────────
Automated test suite specifically verifying the 16 validation scenarios
for UI Redesign & Excel Integration using the real Test.xlsx workbook.

All write operations use temporary non-destructive copies to protect the master workbook.
"""

import os
import shutil
import tempfile
import pytest
import openpyxl

from config import settings
from core.models import (
    EvaluationStatus,
    SaveMode,
    SaveStatus,
    StabilityStatus,
)
from core.exceptions import WorkbookLockedError
from infrastructure.excel_repository import ExcelRepository
from services.excel_service import ExcelService
from services.evaluation_service import EvaluationService
from controllers.measurement_controller import MeasurementController
import main


TEST_XLSX_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "Test.xlsx")
EXPECTED_SHEETS = [
    "ACCURACY STALL",
    "CAM OUTPUT STALL",
    "ACCURACY AUTOPILOT",
    "CAM OUTPUT AUTOPILOT",
    "ACCURACY FLAP",
    "CAM OUTPUT FLAP",
]


@pytest.fixture
def temp_test_workbook():
    """Create a temporary non-destructive copy of Test.xlsx for write tests."""
    tmpdir = tempfile.mkdtemp()
    tmp_path = os.path.join(tmpdir, "Test_Working_Copy.xlsx")
    shutil.copyfile(TEST_XLSX_PATH, tmp_path)
    yield tmp_path
    try:
        shutil.rmtree(tmpdir, ignore_errors=True)
    except Exception:
        pass


# ══════════════════════════════════════════════════════════════════════
#  1. Membuka workbook Test.xlsx
# ══════════════════════════════════════════════════════════════════════
def test_01_open_test_xlsx():
    assert os.path.exists(TEST_XLSX_PATH), f"Master file Test.xlsx must exist at {TEST_XLSX_PATH}"
    wb = openpyxl.load_workbook(TEST_XLSX_PATH, data_only=False)
    assert wb is not None
    assert len(wb.sheetnames) >= 6
    wb.close()


# ══════════════════════════════════════════════════════════════════════
#  2. Memuat keenam worksheet
# ══════════════════════════════════════════════════════════════════════
def test_02_load_six_worksheets():
    wb = openpyxl.load_workbook(TEST_XLSX_PATH, read_only=True)
    sheet_names = wb.sheetnames
    for expected in EXPECTED_SHEETS:
        assert expected in sheet_names, f"Worksheet '{expected}' must be present in Test.xlsx"
    wb.close()


# ══════════════════════════════════════════════════════════════════════
#  3. Memastikan header dan data awal tetap benar
# ══════════════════════════════════════════════════════════════════════
def test_03_headers_and_initial_structure():
    wb = openpyxl.load_workbook(TEST_XLSX_PATH, data_only=False)
    ws = wb["ACCURACY STALL"]

    # Header is at row 9
    assert ws["A9"].value == "TEST POINT"
    assert ws["B9"].value == "INPUT TURNS"
    assert ws["C9"].value == "SHAFT ANGLE"
    assert ws["D9"].value == "SYNCRO ELECT. OUTPUT"
    assert ws["E9"].value == "TOP LIMIT"
    assert ws["F9"].value == "BOTTOM LIMIT"
    assert ws["G9"].value == "RESULT"
    assert ws["H9"].value == "REMARK"
    assert ws["I9"].value == "DIFFERENCE"
    assert ws["J9"].value == "REFERENCE"
    assert ws["K9"].value == "REAL VALUE"

    # Data starts at row 11
    assert ws["A11"].value is not None
    assert ws["D11"].value is not None
    wb.close()


# ══════════════════════════════════════════════════════════════════════
#  4. Mengganti worksheet dan memverifikasi data tabel
# ══════════════════════════════════════════════════════════════════════
def test_04_switch_worksheet_and_verify_table(temp_test_workbook):
    wb = openpyxl.load_workbook(temp_test_workbook)
    repo = ExcelRepository()
    repo.inject(wb, temp_test_workbook)
    svc = ExcelService(repo)

    for sheet in EXPECTED_SHEETS:
        svc.activate_sheet(sheet)
        records = svc.get_live_worksheet_records()
        assert len(records) > 0, f"Sheet '{sheet}' should produce live worksheet records"
        first = records[0]
        assert first["excel_row"] == 11
        assert "test_point" in first
        assert "syncro" in first
        assert "result" not in first or first["measured_angle"] is None or isinstance(first["measured_angle"], (int, float))


# ══════════════════════════════════════════════════════════════════════
#  5. Memilih baris tertentu dan memastikan nomor baris fisik benar
# ══════════════════════════════════════════════════════════════════════
def test_05_select_row_and_verify_physical_mapping(temp_test_workbook):
    wb = openpyxl.load_workbook(temp_test_workbook)
    repo = ExcelRepository()
    repo.inject(wb, temp_test_workbook)
    svc = ExcelService(repo)
    svc.activate_sheet("ACCURACY STALL")

    records = svc.get_live_worksheet_records()
    # Visual index 6 is physical row 17
    row_17_rec = next((r for r in records if r["excel_row"] == 17), None)
    assert row_17_rec is not None
    assert row_17_rec["excel_row"] == 17
    assert row_17_rec["no"] == 7  # 1-indexed count from row 11


# ══════════════════════════════════════════════════════════════════════
#  6. Menyimpan record dan memverifikasi hasil aktual pada workbook
# ══════════════════════════════════════════════════════════════════════
def test_06_save_record_and_verify_workbook_commit(temp_test_workbook):
    wb = openpyxl.load_workbook(temp_test_workbook)
    repo = ExcelRepository()
    repo.inject(wb, temp_test_workbook)
    svc = ExcelService(repo)
    svc.activate_sheet("ACCURACY STALL")

    res = svc.append_measurement(276.71, target_row=11, save_mode="MANUAL")
    assert res.status == SaveStatus.SAVED
    assert res.excel_row == 11

    # Re-read committed file on disk
    wb_verify = openpyxl.load_workbook(temp_test_workbook, data_only=True)
    val = wb_verify["ACCURACY STALL"]["G11"].value
    assert isinstance(val, (int, float))
    assert abs(val - 276.71) < 0.001
    wb_verify.close()


# ══════════════════════════════════════════════════════════════════════
#  7. Mengedit record tanpa mengubah baris lain
# ══════════════════════════════════════════════════════════════════════
def test_07_edit_record_without_disturbing_other_rows(temp_test_workbook):
    wb = openpyxl.load_workbook(temp_test_workbook)
    repo = ExcelRepository()
    repo.inject(wb, temp_test_workbook)
    svc = ExcelService(repo)
    svc.activate_sheet("ACCURACY STALL")

    # Populate rows 11, 12, 13
    svc.append_measurement(111.11, target_row=11)
    svc.append_measurement(222.22, target_row=12)
    svc.append_measurement(333.33, target_row=13)

    # Edit row 12 only
    rec = svc.update_record(12, 999.99)
    assert rec is not None
    assert rec.angle == 999.99

    wb_verify = openpyxl.load_workbook(temp_test_workbook, data_only=True)
    ws = wb_verify["ACCURACY STALL"]
    assert abs(ws["G11"].value - 111.11) < 0.001
    assert abs(ws["G12"].value - 999.99) < 0.001
    assert abs(ws["G13"].value - 333.33) < 0.001
    wb_verify.close()
    wb.close()


# ══════════════════════════════════════════════════════════════════════
#  8. Menjalankan Undo dan memverifikasi perubahan
# ══════════════════════════════════════════════════════════════════════
def test_08_undo_and_verify_changes(temp_test_workbook):
    wb = openpyxl.load_workbook(temp_test_workbook)
    repo = ExcelRepository()
    repo.inject(wb, temp_test_workbook)
    svc = ExcelService(repo)
    svc.activate_sheet("ACCURACY STALL")

    svc.append_measurement(150.00, target_row=11)
    svc.append_measurement(250.00, target_row=12)

    undone = svc.undo_last()
    assert undone is not None
    assert undone.excel_row == 12

    wb_verify = openpyxl.load_workbook(temp_test_workbook, data_only=True)
    ws = wb_verify["ACCURACY STALL"]
    assert abs(ws["G11"].value - 150.00) < 0.001
    assert ws["G12"].value is None
    wb_verify.close()
    wb.close()


# ══════════════════════════════════════════════════════════════════════
#  9. Menguji Auto-Save ON dan Manual Mode
# ══════════════════════════════════════════════════════════════════════
def test_09_auto_save_on_and_manual_mode(temp_test_workbook):
    wb = openpyxl.load_workbook(temp_test_workbook)
    ctrl = MeasurementController()
    ctrl.inject_workbook(wb, temp_test_workbook)
    ctrl.select_sheet("ACCURACY STALL")

    # Set to manual mode
    ctrl.set_save_mode(SaveMode.MANUAL)
    assert ctrl.save_mode == SaveMode.MANUAL
    assert ctrl.auto_save_enabled is False

    # Manual save
    res = ctrl.save_manual(277.00, target_row=11)
    assert res.status == SaveStatus.SAVED
    assert res.excel_row == 11

    # Switch to Auto-Save mode
    ctrl.set_save_mode(SaveMode.AUTO)
    assert ctrl.save_mode == SaveMode.AUTO
    assert ctrl.auto_save_enabled is True


# ══════════════════════════════════════════════════════════════════════
#  10. Memastikan Enter tidak menggandakan record secara tidak sengaja
# ══════════════════════════════════════════════════════════════════════
def test_10_enter_does_not_duplicate_record(temp_test_workbook):
    wb = openpyxl.load_workbook(temp_test_workbook)
    ctrl = MeasurementController()
    ctrl.inject_workbook(wb, temp_test_workbook)
    ctrl.select_sheet("ACCURACY STALL")
    ctrl.set_save_mode(SaveMode.MANUAL)

    # First save succeeds
    res1 = ctrl.save_manual(277.00, target_row=11)
    assert res1.status == SaveStatus.SAVED

    # Rapid second save within debounce window (300ms) should be rejected
    res2 = ctrl.save_manual(277.00, target_row=11)
    assert res2.status == SaveStatus.FAILED
    assert "rapid" in res2.error.lower() or "debounce" in res2.error.lower()


# ══════════════════════════════════════════════════════════════════════
#  11. Memastikan Pause mencegah Auto-Save
# ═══════════════════════════════════════════════════════════════════
def test_11_pause_prevents_autosave(temp_test_workbook):
    wb = openpyxl.load_workbook(temp_test_workbook)
    ctrl = MeasurementController()
    ctrl.inject_workbook(wb, temp_test_workbook)
    ctrl.select_sheet("ACCURACY STALL")
    ctrl.enable_auto_save()

    # Toggle to paused
    paused = ctrl.toggle_auto_save_pause()
    assert paused is True
    assert ctrl.is_auto_save_paused is True

    # Feed angles while paused — auto-save must NOT fire
    saved_calls = []
    ctrl._on_auto_save = lambda r: saved_calls.append(r)

    for _ in range(15):
        ctrl.feed_angle(277.00, is_simulator=True)

    assert len(saved_calls) == 0


# ══════════════════════════════════════════════════════════════════════
#  12. Memverifikasi PASS/FAIL dengan aturan yang digunakan workbook
# ══════════════════════════════════════════════════════════════════════
def test_12_verify_pass_fail_rules():
    eval_svc = EvaluationService(default_tolerance_deg=1.50)
    setpoint = 277.00

    # Within tolerance [275.50, 278.50]
    res_pass = eval_svc.evaluate(277.20, setpoint=setpoint, tolerance=1.50)
    assert res_pass.status == EvaluationStatus.PASS
    assert abs(res_pass.deviation - 0.20) < 0.001

    # Outside tolerance
    res_fail = eval_svc.evaluate(279.00, setpoint=setpoint, tolerance=1.50)
    assert res_fail.status == EvaluationStatus.FAIL
    assert abs(res_fail.deviation - 2.00) < 0.001


# ══════════════════════════════════════════════════════════════════════
#  13. Menguji nilai di batas atas dan bawah toleransi
# ══════════════════════════════════════════════════════════════════════
def test_13_boundary_tolerance_values():
    eval_svc = EvaluationService(default_tolerance_deg=1.50)
    sp = 277.00
    tol = 1.50

    # Exact Top Limit: 277.00 + 1.50 = 278.50 -> PASS
    res_top = eval_svc.evaluate(278.50, setpoint=sp, tolerance=tol)
    assert res_top.status == EvaluationStatus.PASS

    # Just above Top Limit: 278.51 -> FAIL
    res_top_fail = eval_svc.evaluate(278.51, setpoint=sp, tolerance=tol)
    assert res_top_fail.status == EvaluationStatus.FAIL

    # Exact Bottom Limit: 277.00 - 1.50 = 275.50 -> PASS
    res_bot = eval_svc.evaluate(275.50, setpoint=sp, tolerance=tol)
    assert res_bot.status == EvaluationStatus.PASS

    # Just below Bottom Limit: 275.49 -> FAIL
    res_bot_fail = eval_svc.evaluate(275.49, setpoint=sp, tolerance=tol)
    assert res_bot_fail.status == EvaluationStatus.FAIL


# ══════════════════════════════════════════════════════════════════════
#  14. Memastikan formula, format, dan struktur worksheet tidak rusak
# ══════════════════════════════════════════════════════════════════════
def test_14_formulas_and_structure_preserved(temp_test_workbook):
    wb_pre = openpyxl.load_workbook(temp_test_workbook, data_only=False)
    ws_pre = wb_pre["ACCURACY STALL"]
    formula_e11_pre = ws_pre["E11"].value
    formula_f11_pre = ws_pre["F11"].value
    formula_h11_pre = ws_pre["H11"].value
    wb_pre.close()

    # Append data into G11
    repo = ExcelRepository()
    repo.inject(openpyxl.load_workbook(temp_test_workbook), temp_test_workbook)
    svc = ExcelService(repo)
    svc.activate_sheet("ACCURACY STALL")
    svc.append_measurement(276.50, target_row=11)

    # Re-verify formulas remain intact
    wb_post = openpyxl.load_workbook(temp_test_workbook, data_only=False)
    ws_post = wb_post["ACCURACY STALL"]
    assert ws_post["E11"].value == formula_e11_pre
    assert ws_post["F11"].value == formula_f11_pre
    assert ws_post["H11"].value == formula_h11_pre
    assert ws_post["G11"].value == 276.50
    wb_post.close()


# ══════════════════════════════════════════════════════════════════════
#  15. Menguji kondisi file terkunci atau gagal dibaca
# ══════════════════════════════════════════════════════════════════════
def test_15_locked_workbook_handling(temp_test_workbook, monkeypatch):
    wb = openpyxl.load_workbook(temp_test_workbook)
    repo = ExcelRepository()
    repo.inject(wb, temp_test_workbook)
    svc = ExcelService(repo)
    svc.activate_sheet("ACCURACY STALL")

    # Simulate PermissionError on workbook.save and mock Live COM to return False
    def mock_save(*args, **kwargs):
        raise PermissionError("File locked by Microsoft Excel")

    monkeypatch.setattr(wb, "save", mock_save)
    monkeypatch.setattr("infrastructure.excel_repository.LiveExcelAdapter.write_cell_live", staticmethod(lambda *a, **k: False))

    try:
        with pytest.raises(WorkbookLockedError):
            svc.append_measurement(277.00, target_row=11)
    finally:
        wb.close()


# ══════════════════════════════════════════════════════════════════════
#  16. Memastikan GUI tetap responsif dan synchronized dengan Test.xlsx
# ══════════════════════════════════════════════════════════════════════
def test_16_gui_responsiveness_and_summary_sync(temp_test_workbook):
    app = main.FlapSensorApp()
    app.withdraw()

    try:
        wb = openpyxl.load_workbook(temp_test_workbook)
        app._workbook = wb
        app._excel_path = temp_test_workbook
        app._ctrl.inject_workbook(wb, temp_test_workbook)
        app._ctrl.select_sheet("ACCURACY STALL")
        app._active_sheet_name = "ACCURACY STALL"
        app._set_active_row(11)

        # Update telemetry live display
        app._update_live_display(276.85)

        # Verify summary card updated
        assert "276.85" in app._lbl_sum_measured.cget("text")
        assert "ACCURACY STALL" in app._lbl_sum_sheet.cget("text")
        assert "PASS" in app._lbl_sum_eval.cget("text")

        # Refresh live view
        app._refresh_excel_live_view()
        children = app._tree_live.get_children()
        assert len(children) >= 1
    finally:
        app._on_close()
