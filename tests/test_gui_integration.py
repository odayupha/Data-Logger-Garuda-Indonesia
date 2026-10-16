"""
tests/test_gui_integration.py
──────────────────────────────
Integration tests verifying FlapSensorApp properly initializes
with the modular backend architecture and handles basic controller interactions.
"""

import pytest
import openpyxl
import main
from core.models import SaveMode, SaveStatus


@pytest.fixture(scope="module")
def app_instance():
    """Create a headless FlapSensorApp instance and tear down after tests."""
    app = main.FlapSensorApp()
    app.withdraw()
    yield app
    try:
        app._on_close()
    except Exception:
        pass


def test_app_controller_integration(app_instance):
    """Verify controller is attached and initialized."""
    app = app_instance
    assert hasattr(app, "_ctrl")
    assert app._ctrl is not None
    assert app._ctrl.is_workbook_loaded is False


def test_app_workbook_injection(app_instance):
    """Verify injecting a workbook synchronizes controller and UI state."""
    app = app_instance

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TEST_FLAP"
    ws["G11"] = 14.50

    app._workbook = wb
    app._excel_path = "test_gui.xlsx"
    app._ctrl.inject_workbook(wb, "test_gui.xlsx")
    app._ctrl.select_sheet("TEST_FLAP")
    app._active_sheet_name = "TEST_FLAP"

    assert app._ctrl.is_workbook_loaded is True
    assert app._ctrl.active_sheet == "TEST_FLAP"
    assert app._ctrl.scan_next_empty_row() == 12


def test_app_auto_save_toggle(app_instance):
    """Verify toggling auto-save enables/disables controller auto-save."""
    app = app_instance
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TEST_SHEET"

    app._workbook = wb
    app._excel_path = "test_toggle.xlsx"
    app._ctrl.inject_workbook(wb, "test_toggle.xlsx")
    app._active_sheet_name = "TEST_SHEET"

    app._var_auto_save.set(True)
    app._on_auto_save_toggled()
    assert app._ctrl.auto_save_enabled is True

    app._var_auto_save.set(False)
    app._on_auto_save_toggled()
    assert app._ctrl.auto_save_enabled is False


def test_app_load_history_advances_row_when_row_11_filled(app_instance):
    """Verify loading a sheet with row 11 filled automatically targets row 12."""
    app = app_instance
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INSPECTION"
    ws["G11"] = 25.40

    app._workbook = wb
    app._excel_path = "test_adv.xlsx"
    app._ctrl.inject_workbook(wb, "test_adv.xlsx")
    app._active_sheet_name = "INSPECTION"
    app._ctrl.select_sheet("INSPECTION")

    app._load_history_from_excel()

    assert app._current_row == 12
    assert app._var_target_row.get() == "12"
    assert app._ctrl.current_row == 12
    assert len(app._history_data) == 1
    assert app._history_data[0]["excel_row"] == 11
    assert app._history_data[0]["value"] == 25.40


def test_app_restart_row_resets_pointer_and_arms_autosave(app_instance):
    """Verify resetting row 11 resets target pointer and re-arms auto-save."""
    app = app_instance
    app._set_active_row(15)
    assert app._current_row == 15

    app._var_auto_save.set(True)
    app._reset_row_to_default()

    assert app._current_row == 11
    assert app._var_target_row.get() == "11"
    assert app._ctrl.current_row == 11
    assert app._ctrl._auto_save.latch_state.name == "OPEN"


def test_quick_action_bar_cleaned_for_autosave(app_instance):
    """Verify manual save, undo, and export buttons are removed from quick action bar."""
    app = app_instance
    assert hasattr(app, "_btn_edit") and app._btn_edit is not None
    assert hasattr(app, "_btn_restart_bar") and app._btn_restart_bar is not None
    assert getattr(app, "_btn_save", None) is None
    assert getattr(app, "_btn_undo", None) is None
    assert getattr(app, "_btn_export", None) is None


def test_play_sound_beep_callable():
    """Verify play_sound_beep runs without raising exceptions."""
    main.play_sound_beep(1600, 50, is_success=True)
    main.play_sound_beep(450, 50, is_success=False)


def test_auto_save_result_updates_history_and_next_row_pointer(app_instance):
    """Verify auto-save result updates history, highlights row, and advances next target row."""
    app = app_instance
    res = main.SaveResult(
        status=main.SaveStatus.SAVED,
        excel_row=11,
        timestamp_str="12:00:00",
        message="Saved 15.00° → G11",
        angle=15.00,
    )
    app._ctrl.current_row = 12
    app._apply_auto_save_result(res)

    assert app._last_saved_row == 11
    assert app._current_row == 12
    assert app._var_target_row.get() == "12"
    assert "ROW 11 SAVED" in app._history_badge.cget("text")
    assert "NEXT: ROW 12" in app._history_badge.cget("text")
    assert "ROW 12" in app._lbl_action_banner.cget("text")
    # Verify row is recorded with exact 15.00 angle
    assert any(h["excel_row"] == 11 and h["value"] == 15.00 for h in app._history_data)


def test_app_auto_save_pause_toggle(app_instance):
    """Verify pausing auto-save updates controller state, button text, and action banner."""
    app = app_instance
    try:
        # Initially active
        if app._ctrl.is_auto_save_paused:
            app._toggle_auto_save_pause()
        assert app._ctrl.is_auto_save_paused is False
        assert "PAUSE (SPACE)" in app._btn_pause.cget("text")

        # Toggle pause
        app._toggle_auto_save_pause()
        assert app._ctrl.is_auto_save_paused is True
        assert "START (SPACE)" in app._btn_pause.cget("text")
        assert "DI-PAUSE" in app._lbl_action_banner.cget("text")
        assert "PAUSED" in app._lbl_stability_badge.cget("text")

        # Toggle back to start/resume
        app._toggle_auto_save_pause()
        assert app._ctrl.is_auto_save_paused is False
        assert "PAUSE (SPACE)" in app._btn_pause.cget("text")
        assert "AUTO-SAVE AKTIF" in app._lbl_action_banner.cget("text")
    finally:
        if app._ctrl.is_auto_save_paused:
            app._toggle_auto_save_pause()


def test_on_space_pressed_toggles_pause(app_instance):
    """Verify Spacebar press triggers pause toggle when focus is outside text entry fields."""
    app = app_instance
    try:
        if app._ctrl.is_auto_save_paused:
            app._toggle_auto_save_pause()
        assert app._ctrl.is_auto_save_paused is False

        class DummyEvent:
            def __init__(self, widget):
                self.widget = widget

        # Pressing space outside Entry triggers toggle to paused
        dummy_widget = app._btn_pause
        app._on_space_pressed(DummyEvent(dummy_widget))
        assert app._ctrl.is_auto_save_paused is True

        # Pressing space inside an Entry widget does NOT toggle
        # Mock widget with winfo_class returning "Entry"
        class MockEntryWidget:
            def winfo_class(self):
                return "Entry"

        app._on_space_pressed(DummyEvent(MockEntryWidget()))
        assert app._ctrl.is_auto_save_paused is True  # Unchanged!
    finally:
        if app._ctrl.is_auto_save_paused:
            app._toggle_auto_save_pause()


def test_gui_save_mode_segmented_control(app_instance):
    """Verify switching save mode via UI segmented control updates controller and UI badges."""
    app = app_instance
    app._on_save_mode_changed("AUTO-SAVE OFF / MANUAL MODE")
    assert app._ctrl.save_mode == SaveMode.MANUAL
    assert "MANUAL" in app._lbl_mode_badge.cget("text")

    app._on_save_mode_changed("AUTO-SAVE ON")
    assert app._ctrl.save_mode == SaveMode.AUTO
    assert "AUTO-SAVE ENABLED" in app._lbl_mode_badge.cget("text")


def test_gui_enter_pressed_ignored_inside_entry_widget(app_instance):
    """Verify Enter key is ignored when focus is on text Entry field."""
    app = app_instance
    app._live_value = 22.5
    save_called = False

    def mock_save():
        nonlocal save_called
        save_called = True

    orig_save = app._do_manual_save
    app._do_manual_save = mock_save
    try:
        class MockEntryWidget:
            def winfo_class(self):
                return "Entry"

        class DummyEvent:
            def __init__(self, widget):
                self.widget = widget

        app._on_enter_pressed(DummyEvent(MockEntryWidget()))
        assert save_called is False
    finally:
        app._do_manual_save = orig_save


def test_gui_excel_live_view_populates_rows(app_instance):
    """Verify _refresh_excel_live_view populates _tree_live with accurate columns and tags."""
    app = app_instance
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TEST_LIVE_GUI"
    ws["G11"] = 15.20
    ws["G12"] = 30.40

    app._workbook = wb
    app._excel_path = "test_gui_live.xlsx"
    app._ctrl.inject_workbook(wb, "test_gui_live.xlsx")
    app._active_sheet_name = "TEST_LIVE_GUI"
    app._var_target_setpoint.set("15.00")
    app._var_target_tolerance.set("0.50")

    app._refresh_excel_live_view()
    children = app._tree_live.get_children()
    assert len(children) >= 2

    first_vals = app._tree_live.item(children[0], "values")
    assert first_vals[1] == "Row 11"
    assert "15.20" in first_vals[8]
    assert "PASS" in first_vals[9]

    second_vals = app._tree_live.item(children[1], "values")
    assert second_vals[1] == "Row 12"
    assert "30.40" in second_vals[8]
    assert "FAIL" in second_vals[9]







